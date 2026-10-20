#!/usr/bin/env python3
"""
=============================================================================
Premium Technical Report Generator: Word (.docx) Documentation
Script: scripts/build_word_documentation.py
Project: Spatiotemporal CNN-LSTM Framework for Maharashtra Kharif Crop Failure
=============================================================================
Transforms the project documentation into a publication-grade, professionally
styled Microsoft Word (.docx) technical report for B.Tech AI & Data Science.
=============================================================================
"""

import os
import sys
import re
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as patches

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT = Path(__file__).resolve().parent.parent
ART_DIR = ROOT / "art"
SCRATCH = ROOT / "scratch"
SCRATCH.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------------------
# Color Palette Constants (Scientific Editorial / Forestry & Geospatial Theme)
# ---------------------------------------------------------------------------
HEX_PRIMARY = "0F4C3A"       # Deep Forest Green
HEX_SECONDARY = "0F766E"     # Deep Teal
HEX_SLATE_DARK = "1E293B"    # Charcoal / Slate 800
HEX_SLATE_MID = "475569"     # Slate 600
HEX_SLATE_LIGHT = "F1F5F9"   # Slate 100
HEX_ACCENT_SKY = "0284C7"    # Sky Blue
HEX_ACCENT_AMBER = "B45309"  # Amber 700
HEX_BORDER = "CBD5E1"        # Slate 300
HEX_WHITE = "FFFFFF"
HEX_ROW_ALT = "F8FAFC"       # Light Slate Tint

COLOR_PRIMARY = RGBColor(15, 76, 58)       # 0F4C3A
COLOR_SECONDARY = RGBColor(15, 118, 110)   # 0F766E
COLOR_DARK = RGBColor(30, 41, 59)          # 1E293B
COLOR_MUTED = RGBColor(71, 85, 105)        # 475569
COLOR_WHITE = RGBColor(255, 255, 255)

# ---------------------------------------------------------------------------
# 1. XML Formatting Helpers for Word
# ---------------------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def apply_table_borders(table, border_color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_callout(doc, text, title="IMPORTANT NOTICE", style_type="warning"):
    """Creates a beautifully styled editorial callout box with a colored left border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    
    if style_type == "warning":
        bg_hex = "FFFBEB"
        border_hex = "D97706"
        t_col = RGBColor(180, 83, 9)
        b_col = RGBColor(120, 53, 15)
        icon = "⚠️"
    elif style_type == "info":
        bg_hex = "F0F9FF"
        border_hex = "0284C7"
        t_col = RGBColor(3, 105, 161)
        b_col = RGBColor(12, 74, 110)
        icon = "ℹ️"
    elif style_type == "disclosure":
        bg_hex = "FEF2F2"
        border_hex = "E11D48"
        t_col = RGBColor(190, 18, 60)
        b_col = RGBColor(136, 19, 55)
        icon = "🔍"
    else:
        bg_hex = "F0FDF4"
        border_hex = "16A34A"
        t_col = RGBColor(21, 128, 61)
        b_col = RGBColor(20, 83, 45)
        icon = "✅"
        
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_t = p.add_run(f"{icon} {title}: ")
    r_t.bold = True
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(9.5)
    r_t.font.color.rgb = t_col
    
    r_b = p.add_run(text)
    r_b.font.name = "Calibri"
    r_b.font.size = Pt(9.0)
    r_b.font.color.rgb = b_col
    
    # Trailing spacing
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)


# ---------------------------------------------------------------------------
# 2. Render High-Resolution Architectural Vector Diagrams
# ---------------------------------------------------------------------------
def render_diagrams():
    print("[*] Verifying & rendering architectural vector diagrams...")
    
    # Diagram 1: End-to-End System Architecture
    fig, ax = plt.subplots(figsize=(11, 4.4), dpi=300)
    fig.patch.set_facecolor('#0B1310')
    ax.set_facecolor('#0B1310')
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 4.4)
    ax.axis('off')

    boxes = [
        (0.4, 0.6, 2.2, 3.2, "#13231D", "#22C55E", "1. DATA LAYER", 
         "• DES Yields (1966-2024)\n• IMD Gridded Rainfall (0.25°)\n• NOAA CPC ONI SST Index\n• Sentinel-2 L2A (8 Bands)\n• NASA MODIS 8-Day 500m"),
        (3.0, 0.6, 2.2, 3.2, "#13231D", "#38BDF8", "2. FEATURE ENGINE", 
         "• PMFBY Olympic Baselines\n• Yield Deficit (>15%)\n• 6D Crop One-Hot Vectors\n• Dual-Resolution Blender\n• Temporal Normalizers"),
        (5.6, 0.6, 2.2, 3.2, "#13231D", "#EAB308", "3. CNN-LSTM ENGINE", 
         "• 2D Spatial CNN Extractor\n• Recurrent BiLSTM (T=6)\n• Class-Weighted BCE Loss\n• LOSO Cross-Validation\n• Operational PyTorch Checkpoint"),
        (8.2, 0.6, 2.4, 3.2, "#13231D", "#F43F5E", "4. REACT DASHBOARD", 
         "• SVG Choropleth Map Engine\n• Seamless Satellite Terrain\n• Dual-Axis Rainfall-Yield Chart\n• District Agro-Climate Panel\n• Future El Niño Simulator")
    ]

    for x, y, w, h, bg_c, border_c, title, body in boxes:
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.1,rounding_size=0.15",
                                      facecolor=bg_c, edgecolor=border_c, linewidth=2)
        ax.add_patch(rect)
        ax.text(x + w/2, y + h - 0.35, title, color=border_c, fontsize=10.5, fontweight='bold', ha='center', va='center')
        ax.text(x + 0.15, y + h - 0.75, body, color='#E2E8F0', fontsize=8.2, ha='left', va='top', linespacing=1.6)

    for ax_x in [2.7, 5.3, 7.9]:
        arrow = patches.FancyArrowPatch((ax_x, 2.2), (ax_x + 0.2, 2.2), arrowstyle='->',
                                        mutation_scale=18, color='#38BDF8', linewidth=2.5)
        ax.add_patch(arrow)

    d1_path = SCRATCH / "diagram1_system_architecture.png"
    fig.savefig(d1_path, bbox_inches='tight', pad_inches=0.1)
    plt.close(fig)

    # Diagram 2: Spatiotemporal Neural Network Tensor Flow
    fig, ax = plt.subplots(figsize=(11, 4.0), dpi=300)
    fig.patch.set_facecolor('#0B1310')
    ax.set_facecolor('#0B1310')
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 4.0)
    ax.axis('off')

    stages = [
        (0.3, 1.8, 1.8, 1.8, "#1A2E26", "#22C55E", "Sentinel-2 Patch", "[B, 6, 8, 32, 32]\n8 Bands x 6 Windows\n20m GSD (640m Footprint)"),
        (2.5, 1.8, 1.8, 1.8, "#1A2E26", "#22C55E", "Spatial 2D-CNN", "Conv2D + MaxPool\nAdaptiveAvgPool(1,1)\n-> Visual [B, 6, 24]"),
        (4.7, 0.4, 1.8, 1.8, "#152433", "#38BDF8", "Climate & Crop", "[B, 6, 10]\nRainfall + ONI + Cov\n+ 6D Crop One-Hot"),
        (6.9, 1.1, 1.8, 2.2, "#2A2415", "#EAB308", "Recurrent BiLSTM", "Fused Input: [B, 6, 34]\nHidden Size: 24 (Bi)\nOutput: [B, 6, 48]\nFinal Timestep: h_6"),
        (9.1, 1.3, 1.6, 1.8, "#31181F", "#F43F5E", "Risk Output", "Linear(48->24) + Dropout\nLinear(24->1) + Sigmoid\n-> Prob(Failure) in [0, 1]")
    ]

    for x, y, w, h, bg_c, border_c, title, body in stages:
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.1,rounding_size=0.12",
                                      facecolor=bg_c, edgecolor=border_c, linewidth=1.8)
        ax.add_patch(rect)
        ax.text(x + w/2, y + h - 0.3, title, color=border_c, fontsize=9.5, fontweight='bold', ha='center', va='center')
        ax.text(x + w/2, y + (h-0.3)/2, body, color='#E2E8F0', fontsize=7.8, ha='center', va='center', linespacing=1.4)

    ax.add_patch(patches.FancyArrowPatch((2.2, 2.7), (2.4, 2.7), arrowstyle='->', mutation_scale=15, color='#22C55E', linewidth=2))
    ax.add_patch(patches.FancyArrowPatch((4.4, 2.7), (6.8, 2.4), arrowstyle='->', mutation_scale=15, color='#22C55E', linewidth=2))
    ax.add_patch(patches.FancyArrowPatch((6.6, 1.3), (6.8, 1.8), arrowstyle='->', mutation_scale=15, color='#38BDF8', linewidth=2))
    ax.add_patch(patches.FancyArrowPatch((8.8, 2.2), (9.0, 2.2), arrowstyle='->', mutation_scale=15, color='#EAB308', linewidth=2))

    d2_path = SCRATCH / "diagram2_tensor_pipeline.png"
    fig.savefig(d2_path, bbox_inches='tight', pad_inches=0.1)
    plt.close(fig)

    # Diagram 3: PMFBY Olympic 5-Year Baseline Flow
    fig, ax = plt.subplots(figsize=(10, 3.2), dpi=300)
    fig.patch.set_facecolor('#0B1310')
    ax.set_facecolor('#0B1310')
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 3.2)
    ax.axis('off')

    steps = [
        (0.2, 0.5, 2.0, 2.2, "#1E293B", "#38BDF8", "Step 1: History", "Extract preceding\n5 seasons (Y-5..Y-1)\nfor District & Crop\n(Strictly t < Y)"),
        (2.6, 0.5, 2.1, 2.2, "#1E293B", "#38BDF8", "Step 2: Olympic Sort", "Sort yields:\n[y1, y2, y3, y4, y5]\nDrop lowest (y1)\nDrop highest (y5)"),
        (5.1, 0.5, 2.2, 2.2, "#1E293B", "#EAB308", "Step 3: Baseline", "Average middle 3:\nBaseline = (y2+y3+y4)/3\nAnomaly % =\n(Yield_Y - Base)/Base*100"),
        (7.7, 0.5, 2.1, 2.2, "#1E293B", "#F43F5E", "Step 4: Classification", "If Anomaly <= -15%:\n-> y = 1 (Crop Failure)\nIf Anomaly > -15%:\n-> y = 0 (Resilient)")
    ]

    for x, y, w, h, bg_c, border_c, title, body in steps:
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.08,rounding_size=0.12",
                                      facecolor=bg_c, edgecolor=border_c, linewidth=1.8)
        ax.add_patch(rect)
        ax.text(x + w/2, y + h - 0.35, title, color=border_c, fontsize=9.5, fontweight='bold', ha='center', va='center')
        ax.text(x + w/2, y + (h-0.35)/2, body, color='#E2E8F0', fontsize=8.0, ha='center', va='center', linespacing=1.4)

    for ax_x in [2.3, 4.8, 7.4]:
        arrow = patches.FancyArrowPatch((ax_x, 1.6), (ax_x + 0.2, 1.6), arrowstyle='->',
                                        mutation_scale=15, color='#EAB308', linewidth=2)
        ax.add_patch(arrow)

    d3_path = SCRATCH / "diagram3_pmfby_baseline.png"
    fig.savefig(d3_path, bbox_inches='tight', pad_inches=0.1)
    plt.close(fig)

    print("[+] Architectural vector diagrams generated successfully.")
    return d1_path, d2_path, d3_path


# ---------------------------------------------------------------------------
# 3. Master Word Document Assembly Function
# ---------------------------------------------------------------------------
def build_master_word_document(docx_path: Path, d1_path, d2_path, d3_path):
    print(f"\n[*] Assembling Master Word Documentation: {docx_path}...")
    doc = Document()

    # Configure Margins: 0.75 inches for elegant research layout
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document-wide default font
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(10)
    font_normal.color.rgb = COLOR_DARK

    # Helper functions for text elements
    def add_p(text, bold_prefix=None, space_after=5, line_spacing=1.15):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = "Calibri"
            rb.font.color.rgb = COLOR_DARK
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_DARK
        return p

    def add_bullet(text, bold_prefix=None, space_after=3):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = "Calibri"
            rb.font.color.rgb = COLOR_DARK
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_DARK
        return p

    def add_h1(text, space_before=16, space_after=6):
        h = doc.add_heading(level=1)
        h.paragraph_format.space_before = Pt(space_before)
        h.paragraph_format.space_after = Pt(space_after)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(16.5)
        r.font.color.rgb = COLOR_PRIMARY
        return h

    def add_h2(text, space_before=12, space_after=4):
        h = doc.add_heading(level=2)
        h.paragraph_format.space_before = Pt(space_before)
        h.paragraph_format.space_after = Pt(space_after)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(12.5)
        r.font.color.rgb = COLOR_SECONDARY
        return h

    def add_h3(text, space_before=8, space_after=2):
        h = doc.add_heading(level=3)
        h.paragraph_format.space_before = Pt(space_before)
        h.paragraph_format.space_after = Pt(space_after)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = COLOR_DARK
        return h

    def add_figure(img_path: Path, caption: str, width_in=6.2, bold_label=None):
        if not img_path.exists():
            print(f"[!] Warning: Image not found at: {img_path}")
            return
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(str(img_path), width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(10)
        p_cap.paragraph_format.keep_with_next = False
        
        if bold_label:
            rb = p_cap.add_run(bold_label)
            rb.bold = True
            rb.font.name = "Calibri"
            rb.font.size = Pt(8.8)
            rb.font.color.rgb = COLOR_PRIMARY
            
        rc = p_cap.add_run(caption)
        rc.italic = True
        rc.font.name = "Calibri"
        rc.font.size = Pt(8.5)
        rc.font.color.rgb = COLOR_MUTED

    # =======================================================================
    # SECTION 1: COVER PAGE
    # =======================================================================
    print("[*] Generating Executive Cover Page...")
    
    # Top Decorative Banner Table
    top_band = doc.add_table(rows=1, cols=1)
    top_band.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_band = top_band.cell(0, 0)
    set_cell_background(c_band, HEX_PRIMARY)
    set_cell_margins(c_band, top=140, bottom=140, left=200, right=200)
    p_band = c_band.paragraphs[0]
    p_band.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_band = p_band.add_run("B.TECH CAPSTONE RESEARCH REPORT  •  ARTIFICIAL INTELLIGENCE & DATA SCIENCE")
    r_band.bold = True
    r_band.font.name = "Calibri"
    r_band.font.size = Pt(10)
    r_band.font.color.rgb = COLOR_WHITE

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(14)
    p_sp.paragraph_format.space_after = Pt(0)

    # Main Project Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_t1 = p_title.add_run("Spatiotemporal CNN–LSTM Framework for Maharashtra Kharif Crop Failure Prediction")
    r_t1.bold = True
    r_t1.font.name = "Calibri"
    r_t1.font.size = Pt(24)
    r_t1.font.color.rgb = COLOR_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("A Unified Multi-Crop Deep Learning & Geospatial Analytics System for ENSO-Driven Agricultural Shock Early Warning")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12.5)
    r_sub.font.color.rgb = COLOR_SECONDARY
    r_sub.italic = True

    # Cover Hero Visual: Western India Satellite Terrain Reference Map
    cover_hero = ART_DIR / "Western India Satellite Terrain Map.png"
    if cover_hero.exists():
        p_hero = doc.add_paragraph()
        p_hero.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_hero.paragraph_format.space_before = Pt(2)
        p_hero.paragraph_format.space_after = Pt(12)
        run_hero = p_hero.add_run()
        run_hero.add_picture(str(cover_hero), width=Inches(6.4))

    # Executive Metadata Block (Structured 2-column table)
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_borders(meta_table, "CBD5E1")
    
    meta_pairs = [
        ("Target Scope & Geography", "36 Districts of Maharashtra State, Western India (Deccan & Coastal Konkan)"),
        ("Primary Kharif Crops", "6 Crop Conditioning Identity Modes: Rice, Soybean, Cotton, Bajra, Jowar, Tur"),
        ("Ground Truth Benchmark", "PMFBY 5-Year Rolling Olympic Baseline (Yield Deficit Threshold >= 15.0%)"),
        ("Spatiotemporal Deep Learning", "2D Spatial CNN Extractor + Bidirectional LSTM (T=6) + 6D One-Hot Conditioning"),
        ("Evaluation Protocol", "Leave-One-Season-Out (LOSO) Cross-Validation: 98.3% Disaster Recall, 0.987 ROC-AUC")
    ]
    
    for row_idx, (lbl, val) in enumerate(meta_pairs):
        c0 = meta_table.cell(row_idx, 0)
        c1 = meta_table.cell(row_idx, 1)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF" if row_idx % 2 == 0 else "F8FAFC")
        set_cell_margins(c0, top=70, bottom=70, left=120, right=120)
        set_cell_margins(c1, top=70, bottom=70, left=120, right=120)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(lbl)
        r0.bold = True
        r0.font.name = "Calibri"
        r0.font.size = Pt(8.5)
        r0.font.color.rgb = COLOR_PRIMARY
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = COLOR_DARK

    p_bottom = doc.add_paragraph()
    p_bottom.paragraph_format.space_before = Pt(10)
    p_bottom.paragraph_format.space_after = Pt(0)
    r_bot = p_bottom.add_run("Production Release v2.4.0  •  Hugging Face Hub: UsmanSait/elnino-crop-failure-pbl  •  October 2026")
    r_bot.font.name = "Calibri"
    r_bot.font.size = Pt(8.5)
    r_bot.font.color.rgb = COLOR_MUTED
    r_bot.italic = True

    # =======================================================================
    # SECTION 2: BODY WITH UNLINKED HEADERS & FOOTERS
    # =======================================================================
    sec_body = doc.add_section()
    sec_body.header.is_linked_to_previous = False
    sec_body.footer.is_linked_to_previous = False

    # Running Header
    h_para = sec_body.header.paragraphs[0]
    h_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    h_run = h_para.add_run("Spatiotemporal CNN-LSTM Framework — Maharashtra Kharif Crop Failure")
    h_run.font.name = "Calibri"
    h_run.font.size = Pt(8.0)
    h_run.font.color.rgb = COLOR_MUTED

    # Running Footer with Dynamic Page Numbering
    f_para = sec_body.footer.paragraphs[0]
    f_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    f_run = f_para.add_run("B.Tech Capstone Technical Report | AI & Data Science\t\tPage ")
    f_run.font.name = "Calibri"
    f_run.font.size = Pt(8.0)
    f_run.font.color.rgb = COLOR_MUTED
    
    # Insert Page Number Field
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = "PAGE"
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    f_run._r.append(fld1)
    f_run._r.append(instr)
    f_run._r.append(fld2)
    f_run._r.append(fld3)

    # -----------------------------------------------------------------------
    # Table of Contents Block
    # -----------------------------------------------------------------------
    add_h1("Table of Contents", space_before=4)
    toc_items = [
        ("1. Executive Summary & Problem Statement", "Context, Agronomic Vulnerability, Initial Conceptual Blueprint"),
        ("2. End-to-End System Architecture", "Data Layer, Feature Engineering, Neural Network, GIS Dashboard"),
        ("3. Multi-Source Data Ingestion & Provenance", "DES Statistics, IMD Rainfall, NOAA ONI, Sentinel-2 L2A Rasters"),
        ("4. Agricultural Ground Truth & Olympic Baseline", "PMFBY 5-Year Rolling Olympic Average, Mathematical Formulation"),
        ("5. Spatiotemporal CNN–LSTM Neural Network", "2D-CNN Spatial Extractor, BiLSTM Recurrent Modeling, BCE Loss"),
        ("6. Crop Conditioning & Cropland Proxy Disclosure", "6D One-Hot Representation, Regional Canopy Proxy vs Parcel Segmentation"),
        ("7. Leave-One-Season-Out (LOSO) Validation", "Out-of-Season Evaluation (2019-2022), Per-Crop Metrics, Confusion Matrices"),
        ("8. Geospatial Satellite Terrain Pipeline", "Google Earth Engine MODIS Export, High-Res Blending, SVG Alignment"),
        ("9. Frontend Decision-Support Dashboard", "React 18 Architecture, SVG Choropleth Map, Dual-Axis Analytics"),
        ("10. Future Climate Shock Simulator", "Synthetic ENSO SST Warming Shocks (+0.5°C to +2.5°C), Policy Implications"),
        ("11. Deployment, Reproduction & Code Inventory", "Environment Setup, Training Execution, Hugging Face Hub Data Sync")
    ]
    
    t_toc = doc.add_table(rows=len(toc_items), cols=2)
    t_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_borders(t_toc, "E2E8F0")
    for r_idx, (t_chap, t_desc) in enumerate(toc_items):
        c0 = t_toc.cell(r_idx, 0)
        c1 = t_toc.cell(r_idx, 1)
        set_cell_background(c0, "F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
        set_cell_background(c1, "F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
        set_cell_margins(c0, top=50, bottom=50, left=100, right=100)
        set_cell_margins(c1, top=50, bottom=50, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(t_chap)
        r0.bold = True
        r0.font.name = "Calibri"
        r0.font.size = Pt(8.5)
        r0.font.color.rgb = COLOR_PRIMARY
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(t_desc)
        r1.font.name = "Calibri"
        r1.font.size = Pt(8.0)
        r1.font.color.rgb = COLOR_MUTED
        
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -----------------------------------------------------------------------
    # Chapter 1: Executive Summary & Problem Statement
    # -----------------------------------------------------------------------
    add_h1("1. Executive Summary & Problem Statement")
    add_p(
        "Maharashtra is among India's foremost agricultural agrarian economies, producing a significant share "
        "of the nation's pulses, oilseeds, cotton, and staple coarse cereals. However, over 70% of Maharashtra's "
        "cultivated Kharif acreage is completely rainfed and intensely vulnerable to meteorological anomalies "
        "driven by the El Niño Southern Oscillation (ENSO) and the Indian Ocean Dipole (IOD). Anomalous positive sea-surface "
        "temperature warming in the equatorial Pacific (Niño 3.4 index) disrupts the seasonal monsoon trough, delaying onset, "
        "inducing prolonged in-season dry spells, and depleting root-zone soil moisture. This manifests in catastrophic yield "
        "shocks across staple Kharif crops: Rice in the Konkan coastal belt and eastern Vidarbha, Soybean across Marathwada, "
        "Cotton in Vidarbha, and coarse cereals (Bajra and Jowar) in the rain-shadow western plateau."
    )
    add_p(
        "Historically, agricultural disaster management and crop insurance under the Pradhan Mantri Fasal Bima Yojana (PMFBY) "
        "have relied on post-harvest manual Crop Cutting Experiments (CCEs). This legacy approach introduces severe administrative "
        "delays (often 4 to 9 months post-harvest), human measurement errors, and a total inability to execute preventive "
        "in-season agronomic interventions. This capstone engineering research project addresses these critical challenges by "
        "conceiving, training, and deploying a unified spatiotemporal deep learning framework that provides automated, "
        "in-season crop failure early warnings."
    )

    # Summary KPI Table
    add_h3("Table 1.1: Core Engineering & Algorithmic Summary Metrics")
    kpi_headers = ["Metric / Dimension", "Production Specification", "Validation Baseline", "Significance"]
    kpi_data = [
        ["Geographic Extent", "36 Districts of Maharashtra", "DES Harvest Records (1966–2023)", "Complete state-level coverage"],
        ["Crop Multi-Modal Space", "6 Primary Kharif Crops", "One-Hot Conditioned Vector", "Rice, Soybean, Cotton, Bajra, Jowar, Tur"],
        ["Historical Data Depth", "9,078 District-Crop Records", "58 Continuous Seasons", "Authoritative ground truth foundation"],
        ["Satellite Channels", "8 Spectral Bands (10m–20m)", "6 Fortnightly Kharif Windows", "B2, B3, B4, B8, B11, B12, NDVI, NDWI"],
        ["LOSO Failure Recall", "98.3% (59 / 60 Disasters Caught)", "Zero Temporal Leakage Holdouts", "Prioritizes zero false negatives"],
        ["Overall Model ROC-AUC", "0.987 Across 320 Seasons", "Class-Weighted BCE Loss", "Exceptional class separability"],
        ["Dashboard Latency", "< 35 ms Precomputed Serving", "React 18 + Vite GIS Stack", "Interactive policy & farmer decision support"]
    ]
    t_kpi = doc.add_table(rows=len(kpi_data)+1, cols=len(kpi_headers))
    t_kpi.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_borders(t_kpi, "CBD5E1")
    for c_i, h_txt in enumerate(kpi_headers):
        c = t_kpi.cell(0, c_i)
        set_cell_background(c, HEX_PRIMARY)
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h_txt)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(8.2)
        r.font.color.rgb = COLOR_WHITE
        
    for r_i, r_vals in enumerate(kpi_data):
        bg = HEX_ROW_ALT if r_i % 2 == 0 else HEX_WHITE
        for c_i, val in enumerate(r_vals):
            c = t_kpi.cell(r_i+1, c_i)
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)
            p = c.paragraphs[0]
            r = p.add_run(val)
            if c_i == 0:
                r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(8.0)
            r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Conceptual Sketch from art/draw.jpeg
    draw_img = ART_DIR / "draw.jpeg"
    if draw_img.exists():
        add_figure(
            draw_img,
            caption="Initial conceptualization and architectural blueprint sketched during project inception, defining multi-modal data inputs, spatiotemporal feature tensors, Olympic baseline logic, and the dashboard wireframe.",
            width_in=5.2,
            bold_label="Figure 1.1: Project Conceptualization & Data Flow Blueprint. "
        )

    # -----------------------------------------------------------------------
    # Chapter 2: End-to-End System Architecture
    # -----------------------------------------------------------------------
    add_h1("2. End-to-End System Architecture")
    add_p(
        "The project architecture bridges offline geospatial and meteorological remote sensing pipelines with "
        "recurrent neural network modeling and an ultra-responsive web dashboard. As illustrated in Figure 2.1, "
        "the operational system is divided into four cleanly decoupled tiers:"
    )
    add_bullet(
        "Tier 1 — Data Ingestion & Harmonization: Ingests 58 years of DES Maharashtra harvest yields, IMD gridded daily "
        "rainfall, NOAA CPC Oceanic Niño Index (ONI) anomalies, and European Space Agency (ESA) Sentinel-2 Level-2A multi-spectral rasters.",
        bold_prefix="Data Ingestion: "
    )
    add_bullet(
        "Tier 2 — Feature Engineering & Ground Truth Synthesis: Standardizes district nomenclature across 38 administrative "
        "entities, computes rolling PMFBY Olympic baselines, normalizes spatial and climate tensors, and blends dual-resolution satellite basemaps.",
        bold_prefix="Feature Engineering: "
    )
    add_bullet(
        "Tier 3 — Spatiotemporal CNN–LSTM Neural Network: Extracts visual canopy vigor representations via 2D spatial convolutions, "
        "concatenates 10-dimensional climate and crop identity vectors, and models temporal degradation across six fortnightly Kharif observation windows.",
        bold_prefix="Deep Learning: "
    )
    add_bullet(
        "Tier 4 — Interactive GIS Decision Dashboard: Delivers zero-latency policy decision support through an interactive SVG "
        "choropleth map, synchronized dual-axis charts, and an in-silico future El Niño climate shock simulator.",
        bold_prefix="GIS Dashboard: "
    )

    # Figure 2.1: System Architecture Diagram
    add_figure(
        d1_path,
        caption="High-level end-to-end system architecture showing multi-source data ingestion, feature synthesis, the PyTorch deep learning pipeline, and the React GIS dashboard.",
        width_in=6.4,
        bold_label="Figure 2.1: End-to-End System Architecture Diagram. "
    )

    # Figure 2.2: Implementation Schematic from art/implementation.jpeg
    impl_img = ART_DIR / "implementation.jpeg"
    if impl_img.exists():
        add_figure(
            impl_img,
            caption="Detailed engineering implementation schematic showing data storage boundaries, feature normalization modules, PyTorch checkpoint serialization, and dashboard state management.",
            width_in=6.2,
            bold_label="Figure 2.2: System Implementation Schematic. "
        )

    # -----------------------------------------------------------------------
    # Chapter 3: Multi-Source Data Ingestion & Remote Sensing Provenance
    # -----------------------------------------------------------------------
    add_h1("3. Multi-Source Data Ingestion & Remote Sensing Provenance")
    add_p(
        "A foundational strength of this framework is its reliance on authoritative, open-access meteorological, "
        "agronomic, and earth observation records. Table 3.1 catalogs the six core datasets integrated into the pipeline:"
    )

    src_headers = ["Dataset Name", "Source Agency", "Spatial Extent", "Temporal Window", "Role in System"]
    src_data = [
        ["Kharif Harvest Yield Statistics", "DES Maharashtra & ICRISAT", "36 Districts (38 Entities)", "1966 – 2023 (58 Years)", "Ground-truth yields, Olympic baselines & labels"],
        ["IMD Gridded Rainfall (0.25°)", "India Meteorological Dept", "Maharashtra Grid (0.25°)", "2015 – 2024", "Precipitation totals (mm) and departure anomalies (%)"],
        ["Oceanic Niño Index (ONI)", "NOAA Climate Prediction Center", "Niño 3.4 Region (Pacific)", "1950 – 2024", "Summer ENSO SST anomaly teleconnection metric"],
        ["Sentinel-2 L2A Multi-Spectral", "ESA / Copernicus Open Access", "32x32 Patches @ 20m GSD", "2019 – 2022 (6 Windows)", "8-band canopy reflectance, NDVI vigor, NDWI moisture"],
        ["Copernicus DEM (GLO-30)", "European Space Agency (ESA)", "30m Elevation Grid", "Static Elevation Surface", "Topographic shaded relief of Western Ghats & Deccan"],
        ["NASA MODIS Surface Refl.", "NASA LP DAAC (MOD09A1)", "500m Continuous Grid", "8-Day Composites", "Seamless Arabian Sea and regional outer basemap canvas"]
    ]
    t_src = doc.add_table(rows=len(src_data)+1, cols=len(src_headers))
    t_src.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_borders(t_src, "CBD5E1")
    for c_i, h_txt in enumerate(src_headers):
        c = t_src.cell(0, c_i)
        set_cell_background(c, HEX_SECONDARY)
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h_txt)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(8.2)
        r.font.color.rgb = COLOR_WHITE
        
    for r_i, r_vals in enumerate(src_data):
        bg = HEX_ROW_ALT if r_i % 2 == 0 else HEX_WHITE
        for c_i, val in enumerate(r_vals):
            c = t_src.cell(r_i+1, c_i)
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)
            p = c.paragraphs[0]
            r = p.add_run(val)
            if c_i == 0:
                r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(8.0)
            r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h2("Temporal Span Disparity & Alignment Strategy")
    add_p(
        "A critical engineering consideration is the stark temporal disparity between the datasets: while DES harvest statistics "
        "span 58 continuous years (1966–2023), high-resolution Sentinel-2 satellite imagery is only operational from late 2015 onward, "
        "and gridded district rainfall was ingested from 2015 to 2024. The framework resolves this by leveraging the full 58-year historical "
        "depth to establish robust, statistically sound rolling Olympic yield baselines, while training the multi-modal neural network "
        "over the multi-year satellite observation window (2019–2022)."
    )

    add_h2("Sentinel-2 Multi-Spectral Specifications & Cloud Masking")
    add_p(
        "Sentinel-2 Level-2A surface reflectance rasters were extracted at 20m Ground Sample Distance (GSD), yielding 32x32 pixel "
        "patches representing a 640m x 640m ground parcel footprint. Across each season, six continuous half-open temporal windows "
        "track vegetative development: T1 (Jun 01–15, sowing/germination), T2 (Jun 16–30, vegetative emergence), T3 (Jul 01–15, tillering), "
        "T4 (Jul 16–31, peak vegetative growth), T5 (Aug 01–15, flowering/reproductive stage), and T6 (Aug 16–31, grain filling)."
    )
    add_p(
        "Each patch contains eight spectral channels: Blue (B2, 490 nm), Green (B3, 560 nm), Red (B4, 665 nm), Near-Infrared (B8, 842 nm), "
        "Short-Wave Infrared 1 (B11, 1610 nm), Short-Wave Infrared 2 (B12, 2190 nm), Normalized Difference Vegetation Index (NDVI), and "
        "Normalized Difference Water Index (NDWI). Cloud and shadow contamination is masked using the Sentinel-2 Scene Classification Layer (SCL), "
        "and a valid pixel coverage metric C_t is recorded per window."
    )

    sat_img = ART_DIR / "Sattelite.png"
    if sat_img.exists():
        add_figure(
            sat_img,
            caption="Multi-spectral satellite patch composition showing the 8 spectral channels across consecutive fortnightly observation windows, illustrating vegetative green-up and moisture dynamics.",
            width_in=6.2,
            bold_label="Figure 3.1: Sentinel-2 Multi-Spectral Band Composition. "
        )

    # -----------------------------------------------------------------------
    # Chapter 4: Agricultural Ground Truth & PMFBY Olympic Baseline
    # -----------------------------------------------------------------------
    add_h1("4. Agricultural Ground Truth & PMFBY Olympic Baseline")
    add_p(
        "Defining agricultural failure through a fixed absolute yield threshold (e.g., < 1,000 kg/ha) is statistically invalid "
        "due to substantial agro-climatic disparities across districts (e.g., irrigated Pune vs rainfed Gadchiroli) and long-term "
        "technological yield creep from improved hybrid seeds and fertilizer adoption. To establish ground truth aligned with national "
        "crop insurance jurisprudence, this project implements the official Pradhan Mantri Fasal Bima Yojana (PMFBY) 5-year rolling "
        "Olympic baseline methodology."
    )

    add_figure(
        d3_path,
        caption="Four-stage PMFBY 5-year rolling Olympic baseline workflow and binary failure classification threshold.",
        width_in=6.2,
        bold_label="Figure 4.1: PMFBY Rolling Olympic Baseline Logic. "
    )

    add_h2("Mathematical Formulation & Derivation")
    add_p("For a given district D, crop C, and target season Y, the algorithm executes four formal mathematical steps:")
    add_p("1. Historical Candidate Set Extraction: Collect preceding valid historical harvest yields over a 7-year lookback window:")
    add_p("       Y_prior = { Yield_(D,C,t) | Y - 7 <= t < Y  and  Yield_(D,C,t) > 0 }", bold_prefix="   Lookback: ")
    add_p("   Strict temporal integrity is enforced: no concurrent or future data is ever included in the baseline calculation.")
    
    add_p("2. Olympic Sorting: When at least five valid prior seasons exist, the most recent five yields are sorted in ascending order:")
    add_p("       y_(1) <= y_(2) <= y_(3) <= y_(4) <= y_(5)", bold_prefix="   Sorted Sequence: ")
    add_p("   The extreme lowest value y_(1) (potential catastrophic outlier) and extreme highest value y_(5) (bumper harvest) are dropped.")

    add_p("3. Baseline Calculation: The Olympic baseline yield is the arithmetic mean of the middle three normal seasons:")
    add_p("       Baseline_(D,C,Y) = ( y_(2) + y_(3) + y_(4) ) / 3", bold_prefix="   Olympic Mean: ")

    add_p("4. Yield Anomaly & Failure Labeling: The percentage departure from normal historical expectation is computed as:")
    add_p("       Anomaly_% = [ ( Yield_(D,C,Y) - Baseline_(D,C,Y) ) / Baseline_(D,C,Y) ] * 100%", bold_prefix="   Yield Anomaly: ")
    add_p("       Label y_(D,C,Y) = 1 (Crop Failure) if Anomaly_% <= -15.0%, else 0 (Resilient)", bold_prefix="   Binary Target: ")

    add_callout(
        doc,
        "Numerical Demonstration (Gondia Rice 2019): Preceding 5 valid seasons: 2014 (1,740 kg/ha), 2015 (1,480 kg/ha), "
        "2016 (1,890 kg/ha), 2017 (1,620 kg/ha), 2018 (1,550 kg/ha). Sorted: [1480, 1550, 1620, 1740, 1890]. Dropping extremes "
        "(1480, 1890) leaves [1550, 1620, 1740]. Baseline = (1550+1620+1740)/3 = 1,636.7 kg/ha. In 2019, actual harvest yield was "
        "1,295 kg/ha. Anomaly = ((1295 - 1636.7) / 1636.7) * 100% = -20.88%. Because -20.88% <= -15.0%, the season is classified "
        "as a true agronomic Crop Failure (y = 1).",
        title="PMFBY NUMERICAL CASE STUDY",
        style_type="info"
    )

    # -----------------------------------------------------------------------
    # Chapter 5: Spatiotemporal CNN–LSTM Neural Network Architecture
    # -----------------------------------------------------------------------
    add_h1("5. Spatiotemporal CNN–LSTM Neural Network Architecture")
    add_p(
        "To simultaneously capture spatial vegetative canopy vigor and temporal degradation across the growing season, "
        "we designed a dual-branch neural network combining 2D spatial convolutions with a bidirectional recurrent LSTM. "
        "Figure 5.1 illustrates the complete tensor flow and layer hierarchy:"
    )

    add_figure(
        d2_path,
        caption="Spatiotemporal deep learning tensor pipeline showing 2D-CNN feature extraction, climate and crop-conditioning fusion, BiLSTM recurrence, and classification head.",
        width_in=6.4,
        bold_label="Figure 5.1: Neural Network Architecture & Tensor Flow. "
    )

    add_h2("Layer Tensor Shapes & Flow Specifications")
    
    layer_headers = ["Stage / Layer", "Input Tensor Shape", "Operations & Hyperparameters", "Output Tensor Shape"]
    layer_data = [
        ["Satellite Patch Input", "[Batch, 6, 8, 32, 32]", "8 Multi-spectral Channels x 6 Windows", "[Batch*6, 8, 32, 32]"],
        ["Spatial CNN Block 1", "[Batch*6, 8, 32, 32]", "Conv2D(8->16, k=3, p=1) + BN + ReLU + MaxPool(2)", "[Batch*6, 16, 16, 16]"],
        ["Spatial CNN Block 2", "[Batch*6, 16, 16, 16]", "Conv2D(16->24, k=3, p=1) + BN + ReLU + MaxPool(2)", "[Batch*6, 24, 8, 8]"],
        ["Spatial Adaptive Pool", "[Batch*6, 24, 8, 8]", "AdaptiveAvgPool2d((1,1)) + Flatten()", "[Batch, 6, 24]"],
        ["Climate & Conditioning", "[Batch, 6, 10]", "Rainfall (norm) + Departure + ONI + Cov + 6D One-Hot", "[Batch, 6, 10]"],
        ["Sequential Concat", "Spatial [24] + Clim [10]", "torch.cat([spatial_feat, climate_seq], dim=-1)", "[Batch, 6, 34]"],
        ["Recurrent BiLSTM", "[Batch, 6, 34]", "1 Layer, Bidirectional, Hidden Dim = 24", "[Batch, 6, 48]"],
        ["Temporal Pooling", "[Batch, 6, 48]", "Final Timestep State Extraction h_6 = lstm_out[:, -1, :]", "[Batch, 48]"],
        ["Classification Head", "[Batch, 48]", "Linear(48->24) + ReLU + Dropout(0.25) -> Linear(24->1)", "Logits [Batch, 1]"],
        ["Loss Optimization", "Logits & Target y", "BCEWithLogitsLoss(pos_weight = N_neg / N_pos ≈ 4.33)", "Scalar Loss"]
    ]
    t_lay = doc.add_table(rows=len(layer_data)+1, cols=len(layer_headers))
    t_lay.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_borders(t_lay, "CBD5E1")
    for c_i, h_txt in enumerate(layer_headers):
        c = t_lay.cell(0, c_i)
        set_cell_background(c, HEX_PRIMARY)
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h_txt)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(8.2)
        r.font.color.rgb = COLOR_WHITE
        
    for r_i, r_vals in enumerate(layer_data):
        bg = HEX_ROW_ALT if r_i % 2 == 0 else HEX_WHITE
        for c_i, val in enumerate(r_vals):
            c = t_lay.cell(r_i+1, c_i)
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)
            p = c.paragraphs[0]
            r = p.add_run(val)
            if c_i == 0:
                r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(8.0)
            r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h2("Class-Weighted Loss Optimization Rationale")
    add_p(
        "In agricultural disaster warning, the societal and economic cost of a False Negative (failing to alert authorities "
        "to an impending famine or crop collapse) is vastly higher than that of a False Positive (a precautionary advisory). "
        "To reflect this asymmetry, the loss function configures a positive class re-weighting factor: pos_weight = N_neg / N_pos ≈ 4.33. "
        "This penalizes missed failures heavily, forcing the optimizer to achieve near-perfect recall (98.3%)."
    )

    # -----------------------------------------------------------------------
    # Chapter 6: Crop Conditioning & Cropland Proxy Disclosure
    # -----------------------------------------------------------------------
    add_h1("6. Crop Conditioning Mechanism & Regional Cropland Proxy Disclosure")
    add_p(
        "Inside train_all_kharif_crops_cnn_lstm.py, multi-crop prediction is unified into a single deep learning model "
        "through a 6-dimensional one-hot identity vector: Rice [1,0,0,0,0,0], Soybean [0,1,0,0,0,0], Cotton [0,0,1,0,0,0], "
        "Bajra [0,0,0,1,0,0], Jowar [0,0,0,0,1,0], and Tur [0,0,0,0,0,1]. This crop identity vector is appended to the normalized "
        "climate indicators at every timestep of the BiLSTM, enabling the network to learn crop-specific physiological thresholds "
        "under identical meteorological conditions."
    )

    add_callout(
        doc,
        "Critical Agronomic & Remote Sensing Caveat: In the current dataset, Sentinel-2 multi-spectral patches were sampled across "
        "agricultural cropland parcels within 20 canonical districts. The multi-crop neural network pairs these regional cropland patches "
        "with each crop's one-hot identity vector and official DES yield label. Consequently, the satellite branch operates as a "
        "REGIONAL VEGETATIVE CANOPY AND MOISTURE PROXY (reflecting district-wide vegetative vigor via NDVI and moisture deficit via NDWI), "
        "while the crop conditioning vector modulates vulnerability. The model does NOT view field-level segmented crop polygons "
        "(e.g., exclusively delineated cotton or soybean boundaries). This design preserves scalability across all 36 districts but should "
        "be transparently disclosed in scientific publications.",
        title="ACADEMIC & REMOTE SENSING METHODOLOGY DISCLOSURE",
        style_type="disclosure"
    )

    # -----------------------------------------------------------------------
    # Chapter 7: Leave-One-Season-Out (LOSO) Validation
    # -----------------------------------------------------------------------
    add_h1("7. Leave-One-Season-Out (LOSO) Validation & Evaluation Metrics")
    add_p(
        "Standard random k-fold cross-validation is fundamentally flawed for climate-agricultural modeling because temporal autocorrelation "
        "causes severe data leakage (the model memorizes weather patterns from neighboring districts in the same year). To guarantee strict "
        "operational out-of-season generalizability, we evaluated the framework using Leave-One-Season-Out (LOSO) cross-validation: "
        "in each fold, one complete Kharif season (e.g., 2019) is completely sequestered as the test set, while the model is trained exclusively "
        "on the remaining seasons. Normalization statistics (means and standard deviations) were fitted strictly on training folds."
    )

    add_h2("Per-Crop Operational Performance Table (320 District-Crop Seasons)")
    
    m_headers = ["Crop", "Records", "Failures", "Rate", "TN", "FP", "FN", "TP", "Precision", "Recall", "F1", "BalAcc", "ROC-AUC"]
    m_data = [
        ["Bajra", "37", "5", "13.5%", "32", "0", "0", "5", "100.0%", "100.0%", "1.000", "1.000", "1.000"],
        ["Cotton", "49", "8", "16.3%", "40", "1", "0", "8", "88.9%", "100.0%", "0.941", "0.988", "1.000"],
        ["Jowar", "49", "16", "32.7%", "27", "6", "0", "16", "72.7%", "100.0%", "0.842", "0.909", "0.992"],
        ["Rice", "63", "8", "12.7%", "51", "4", "0", "8", "66.7%", "100.0%", "0.800", "0.964", "0.968"],
        ["Soybean", "54", "14", "25.9%", "36", "4", "0", "14", "77.8%", "100.0%", "0.875", "0.950", "0.984"],
        ["Tur", "68", "9", "13.2%", "58", "1", "1", "8", "88.9%", "88.9%", "0.889", "0.936", "0.985"],
        ["Overall", "320", "60", "18.8%", "244", "16", "1", "59", "78.7%", "98.3%", "0.874", "0.961", "0.987"]
    ]
    t_met = doc.add_table(rows=len(m_data)+1, cols=len(m_headers))
    t_met.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_borders(t_met, "CBD5E1")
    for c_i, h_txt in enumerate(m_headers):
        c = t_met.cell(0, c_i)
        set_cell_background(c, HEX_SLATE_DARK)
        set_cell_margins(c, top=70, bottom=70, left=70, right=70)
        p = c.paragraphs[0]
        r = p.add_run(h_txt)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(7.8)
        r.font.color.rgb = COLOR_WHITE
        
    for r_i, r_vals in enumerate(m_data):
        bg = "E2E8F0" if r_i == len(m_data) - 1 else (HEX_ROW_ALT if r_i % 2 == 0 else HEX_WHITE)
        for c_i, val in enumerate(r_vals):
            c = t_met.cell(r_i+1, c_i)
            set_cell_background(c, bg)
            set_cell_margins(c, top=55, bottom=55, left=70, right=70)
            p = c.paragraphs[0]
            r = p.add_run(val)
            if r_i == len(m_data) - 1 or c_i in [0, 8, 9, 10, 12]:
                r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(7.5)
            r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("Analysis of Key Findings")
    add_bullet(
        "Superb Disaster Recall (98.3%): Across all 320 district-crop seasons evaluated, the model correctly identified 59 out of 60 "
        "historical crop failure events, missing only a single disaster (FN = 1 in Tur 2020).",
        bold_prefix="Recall: "
    )
    add_bullet(
        "Controlled False Alarms: Only 16 false positives were recorded across 260 resilient seasons, yielding an overall precision of 78.7%. "
        "This false alarm trade-off is optimal for agrarian disaster preparedness.",
        bold_prefix="Precision: "
    )
    add_bullet(
        "Comparison with Baseline Models: The CNN-LSTM significantly outperforms both naive Majority Class guessing (Recall: 0.0%, F1: 0.000) "
        "and Climate-Only Logistic Regression (Recall: 76.7%, F1: 0.621, ROC-AUC: 0.833), demonstrating that the spatiotemporal satellite "
        "vision branch provides vital canopy moisture signals that meteorological data alone misses.",
        bold_prefix="Baselines: "
    )

    # -----------------------------------------------------------------------
    # Chapter 8: Geospatial Satellite Terrain & Earth Engine Mapping Engine
    # -----------------------------------------------------------------------
    add_h1("8. Geospatial Satellite Terrain & Earth Engine Mapping Engine")
    add_p(
        "A major aesthetic and technical breakthrough in this project is the construction of a seamless, authentic Earth observation "
        "terrain basemap. In standard GIS implementations, high-resolution Digital Elevation Models (such as Copernicus DEM GLO-30) "
        "suffer from severe 'stepping dropouts' along coastal boundaries because land-only DEMs abruptly terminate at the coastline, "
        "producing artificial black voids in the Arabian Sea."
    )
    add_p(
        "To solve this problem, we engineered a three-stage dual-resolution compositing pipeline:"
    )
    add_bullet(
        "Stage 1 (Earth Engine MODIS Export): Using Google Earth Engine (export_satellite_terrain_background.js), we composite NASA MODIS "
        "8-day surface reflectance (MOD09A1 500m) across a broad subcontinental bounding box [62.0°E, 8.0°N, 92.0°E, 30.0°N]. Because MODIS "
        "continuously images both ocean and land, it provides a gapless regional canvas.",
        bold_prefix="GEE MODIS Export: "
    )
    add_bullet(
        "Stage 2 (Dual-Resolution Alpha Blending): In scripts/combine_wide_satellite_terrain.py, the regional 500m MODIS canvas is blended "
        "with the high-resolution Maharashtra GeoTIFF (10m Sentinel-2 + Copernicus DEM) using a 20-pixel Gaussian-blurred alpha transition mask "
        "and color-ratio balancing. A post-export fill pass replaces any remaining southwest corner nodata with deep oceanic RGB(0, 11, 21).",
        bold_prefix="Python Blending: "
    )
    add_bullet(
        "Stage 3 (Precision SVG Coordinate Transform): The resulting WebP raster is mapped to the React SVG viewBox (0, 0, 860, 713) using "
        "affine transformation coordinates: x = -1091.8 px, y = -859.9 px, width = 3089.8 px, height = 2394.6 px. This guarantees zero black "
        "margins even when users zoom out to 75% or pan along the Konkan coastline.",
        bold_prefix="SVG Alignment: "
    )

    # Figure 8.1: Earth Engine Code Editor
    gee_img = ART_DIR / "Earth engine.png"
    if gee_img.exists():
        add_figure(
            gee_img,
            caption="Google Earth Engine JavaScript code editor running the automated satellite extraction pipeline, showing task queues, multi-spectral band compositing, and geospatial asset exports.",
            width_in=6.2,
            bold_label="Figure 8.1: Google Earth Engine Extraction Pipeline. "
        )

    # Figure 8.2: Master Satellite Basemap
    master_basemap = ROOT / "public" / "images" / "maharashtra_satellite_terrain_sentinel2_dem.png"
    if master_basemap.exists():
        add_figure(
            master_basemap,
            caption="Final production dual-resolution satellite terrain basemap blending the high-resolution Maharashtra core with wide MODIS Arabian Sea and Indian subcontinent coverage.",
            width_in=6.2,
            bold_label="Figure 8.2: Production Dual-Resolution Satellite Terrain Basemap. "
        )

    # -----------------------------------------------------------------------
    # Chapter 9: Frontend Decision-Support Dashboard Architecture
    # -----------------------------------------------------------------------
    add_h1("9. Frontend Decision-Support Dashboard & GIS Interface")
    add_p(
        "The web application was engineered as a production-grade GIS decision-support tool using React 18, Vite, TypeScript, and "
        "TailwindCSS. It delivers sub-35ms interactions by ingesting precomputed model predictions alongside live dynamic filtering."
    )

    # Figure 9.1: UI Layout Wireframe from art/layout rough.png
    wireframe_img = ART_DIR / "layout rough.png"
    if wireframe_img.exists():
        add_figure(
            wireframe_img,
            caption="UI/UX wireframe sketch drafted during dashboard design, outlining sidebar navigation, KPI summary cards, interactive choropleth map, and dual-axis chart panel.",
            width_in=5.8,
            bold_label="Figure 9.1: Dashboard UI/UX Blueprint Wireframe. "
        )

    add_h2("Core Frontend Component Modular Hierarchy")
    add_bullet(
        "DashboardContext.tsx: Centralized React Context coordinating global state: selectedYear (2019–2022), selectedDistrictId, "
        "selectedCrop (Rice, Soybean, Cotton, Bajra, Jowar, Tur), and all 320 precomputed CNN-LSTM district-crop inferences.",
        bold_prefix="Global State: "
    )
    add_bullet(
        "DistrictChoropleth.tsx: Interactive SVG map engine rendering 36 Maharashtra district polygons with smooth drag-panning, wheel-zooming "
        "(clamped 75% to 250%), dynamic risk categorization (Severe Risk, Moderate Risk, Resilient, No Data), and real-time satellite terrain alignment.",
        bold_prefix="Choropleth Map: "
    )
    add_bullet(
        "RainfallYieldChart.tsx: Synchronized dual-axis chart contrasting IMD rainfall against historical harvest yields, with vertical indicators "
        "marking El Niño drought years and synchronized climate index docks.",
        bold_prefix="Dual-Axis Chart: "
    )
    add_bullet(
        "DistrictDetailPanel.tsx: Side-docked agronomic telemetry panel presenting district soil classification, Kharif rainfall departure, "
        "satellite NDVI canopy vigor, and PMFBY insurance claim advisory recommendations.",
        bold_prefix="District Panel: "
    )

    # Figure 9.2: Live Production Dashboard Screenshot (art/Screenshot 2026-10-21 033114.png)
    live_dash = ART_DIR / "Screenshot 2026-10-21 033114.png"
    if live_dash.exists():
        add_figure(
            live_dash,
            caption="Live production React 18 GIS dashboard showing the seamless satellite terrain basemap, district failure risk classification choropleth, navigation sidebar, and summary KPI telemetry.",
            width_in=6.4,
            bold_label="Figure 9.2: Production React GIS Decision Dashboard. "
        )

    # Figure 9.3: KPI Summary Cards & Dual-Axis Chart (art/Screenshot 2026-10-21 020442.png)
    kpi_chart_img = ART_DIR / "Screenshot 2026-10-21 020442.png"
    if kpi_chart_img.exists():
        add_figure(
            kpi_chart_img,
            caption="Close-up view of the 6 KPI summary cards, synchronized dual-axis rainfall-yield chart with historical drought markers, and district detail panel.",
            width_in=6.2,
            bold_label="Figure 9.3: KPI Summary Cards & Dual-Axis Rainfall-Yield Analytics. "
        )

    # Figure 9.4: Interactive District Risk Tooltip (art/Screenshot 2026-10-21 014803.png)
    tooltip_img = ART_DIR / "Screenshot 2026-10-21 014803.png"
    if tooltip_img.exists():
        add_figure(
            tooltip_img,
            caption="Interactive district inspection view demonstrating real-time risk hovering, tooltip telemetry display, and satellite terrain alignment over Western Maharashtra.",
            width_in=6.2,
            bold_label="Figure 9.4: Interactive District Inspection & Risk Tooltip. "
        )

    # -----------------------------------------------------------------------
    # Chapter 10: Future Climate Shock Simulator & Agronomic Implications
    # -----------------------------------------------------------------------
    add_h1("10. Future Climate Shock Simulator & Agronomic Policy Implications")
    add_p(
        "To transition the dashboard from a purely retrospective diagnostic tool into an active forward-looking policy instrument, "
        "we implemented the Future El Niño Scenario Simulator (FutureElNinoPage.tsx). Users can inject synthetic sea-surface temperature "
        "warming shocks (+0.5°C to +2.5°C Niño 3.4 anomalies) and simulate how drought vulnerability cascades across individual crops and districts."
    )

    # Figure 10.1: Future El Niño Simulator (art/Screenshot 2026-10-21 020537.png)
    sim_img = ART_DIR / "Screenshot 2026-10-21 020537.png"
    if sim_img.exists():
        add_figure(
            sim_img,
            caption="Future El Niño scenario simulator modeling the probability of crop failure under synthetic ENSO SST warming shocks (+0.5°C to +2.5°C) across Maharashtra districts.",
            width_in=6.2,
            bold_label="Figure 10.1: Future El Niño Climate Shock Simulation Interface. "
        )

    add_h2("Crop Vulnerability Hierarchy & Agronomic Recommendations")
    add_p(
        "Our empirical and simulation findings reveal a clear physiological vulnerability hierarchy across Kharif crops:"
    )
    add_bullet(
        "Rice (Konkan & Vidarbha): Highest sensitivity to dry spells during tillering and panicle initiation. A 20% rainfall deficit during "
        "July triggers near-certain crop collapse (failure probability > 85%). Policy recommendation: mandate early nursery transplanting and "
        "promote Alternate Wetting and Drying (AWD) water management.",
        bold_prefix="1. Rice: "
    )
    add_bullet(
        "Soybean (Marathwada): Highly vulnerable to pod-filling moisture stress in August. Sensitive to both drought and waterlogging. "
        "Policy recommendation: expand Broad Bed Furrow (BBF) planting to manage erratic monsoon dry spells.",
        bold_prefix="2. Soybean: "
    )
    add_bullet(
        "Cotton (Vidarbha): Long duration crop vulnerable to extended dry spells causing square and boll shedding. Policy recommendation: "
        "promote drip irrigation subsidies and short-duration varieties.",
        bold_prefix="3. Cotton: "
    )
    add_bullet(
        "Coarse Cereals (Bajra & Jowar): Exceptional drought hardiness. Even under moderate El Niño conditions (+1.0°C ONI), Bajra maintains "
        "resilience in 86.5% of district seasons. Policy recommendation: promote millets as climate-resilient alternative crops in drought-prone talukas.",
        bold_prefix="4. Bajra & Jowar: "
    )

    # -----------------------------------------------------------------------
    # Chapter 11: Deployment, Reproduction & Operational Code Inventory
    # -----------------------------------------------------------------------
    add_h1("11. Deployment, Reproduction & Code Inventory")
    add_p("To ensure 100% academic reproducibility, the complete project can be executed and retrained through standard CLI commands:")
    
    add_h2("Step-by-Step Reproduction Guide")
    add_p("1. Clone Repository & Install Dependencies:", bold_prefix="Local Setup: ")
    add_p("       git clone https://github.com/usmanssait/ElninoPBLProject.git\n       cd ElninoPBLProject\n       npm install")
    
    add_p("2. Launch Interactive GIS Dashboard:", bold_prefix="Start Web App: ")
    add_p("       npm run dev\n       # Dashboard executes locally at http://localhost:3000")

    add_p("3. Retrain Multi-Crop CNN–LSTM Deep Learning Model:", bold_prefix="Model Retraining: ")
    add_p("       & 'C:\\Users\\USMAN\\AppData\\Local\\Python\\bin\\python.exe' CNN_LSTM_Training/CropFailureProject/train_all_kharif_crops_cnn_lstm.py")

    add_p("4. Re-generate Dual-Resolution Satellite Terrain Basemap:", bold_prefix="Satellite Pipeline: ")
    add_p("       & 'C:\\Users\\USMAN\\AppData\\Local\\Python\\bin\\python.exe' scripts/combine_wide_satellite_terrain.py -i data/regional_seamless_satellite_terrain_500m.tif")

    add_p("5. Synchronize Large Datasets & Model Checkpoints with Hugging Face Hub:", bold_prefix="HF Hub Sync: ")
    add_p("       & 'C:\\Users\\USMAN\\AppData\\Local\\Python\\bin\\python.exe' hf_sync.py --repo-id UsmanSait/elnino-crop-failure-pbl --repo-type dataset --token <HF_TOKEN>")

    add_h2("Repository Code & Data Inventory")
    inv_headers = ["File / Directory", "Component Type", "Description & Operational Function"]
    inv_data = [
        ["src/context/DashboardContext.tsx", "Frontend State", "Central React Context for year, crop, district, and precomputed inferences"],
        ["src/components/DistrictChoropleth.tsx", "GIS Map Engine", "Interactive SVG choropleth with zoom/pan and satellite terrain alignment"],
        ["src/components/RainfallYieldChart.tsx", "Analytics View", "Synchronized dual-axis chart for rainfall, yields, and drought indicators"],
        ["src/pages/FutureElNinoPage.tsx", "Scenario Simulator", "Policy simulator modeling SST warming shocks (+0.5°C to +2.5°C)"],
        ["CNN_LSTM_Training/.../train_all_kharif_crops_cnn_lstm.py", "Deep Learning", "PyTorch 2D-CNN + BiLSTM multi-crop training and LOSO evaluation engine"],
        ["scripts/build_standardized_kharif_series.py", "Data Engine", "Standardizes 58 years of DES harvest records & computes PMFBY baselines"],
        ["scripts/extract_sentinel2_cropland.js", "Earth Engine", "GEE JavaScript script exporting 8-band fortnightly Sentinel-2 patches"],
        ["scripts/combine_wide_satellite_terrain.py", "Geospatial", "Blends 500m MODIS canvas with 10m Sentinel-2 + DEM core basemap"],
        ["data/standardized_maharashtra_kharif_all_crops_1966_2024.csv", "Data Asset", "9,078 records across 38 district entities and 6 Kharif crops"],
        ["data/crop_failure_predictions.csv", "Predictions", "320 precomputed district-crop predictions ingested by React dashboard"]
    ]
    t_inv = doc.add_table(rows=len(inv_data)+1, cols=len(inv_headers))
    t_inv.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_borders(t_inv, "CBD5E1")
    for c_i, h_txt in enumerate(inv_headers):
        c = t_inv.cell(0, c_i)
        set_cell_background(c, HEX_PRIMARY)
        set_cell_margins(c, top=70, bottom=70, left=90, right=90)
        p = c.paragraphs[0]
        r = p.add_run(h_txt)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(8.0)
        r.font.color.rgb = COLOR_WHITE
        
    for r_i, r_vals in enumerate(inv_data):
        bg = HEX_ROW_ALT if r_i % 2 == 0 else HEX_WHITE
        for c_i, val in enumerate(r_vals):
            c = t_inv.cell(r_i+1, c_i)
            set_cell_background(c, bg)
            set_cell_margins(c, top=50, bottom=50, left=90, right=90)
            p = c.paragraphs[0]
            r = p.add_run(val)
            if c_i == 0:
                r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(7.8)
            r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Save Document
    doc.save(str(docx_path))
    print(f"[+] Master Word documentation successfully generated at: {docx_path}")
    print(f"    File size: {docx_path.stat().st_size / (1024*1024):.2f} MB")


def main():
    d1, d2, d3 = render_diagrams()
    docx_out = ROOT / "FULL_PROJECT_DOCUMENTATION.docx"
    build_master_word_document(docx_out, d1, d2, d3)

if __name__ == '__main__':
    main()
